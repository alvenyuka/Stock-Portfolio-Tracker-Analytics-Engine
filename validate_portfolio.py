"""Re-derive the workbook's figures from its raw inputs, outside the spreadsheet.

The Validation tab reports its own checks passing. Those checks are computed by
the same workbook whose formulas are in question, so they carry no independent
weight. This script is the outside opinion.

It reads only inputs:

* the Ledger: date, ticker, buy or sell, price, units and amount of every trade;
* each holding's current price, and the Stocks data type betas;
* the 12 months of daily closes on the Price History sheet, and the SPY closes;
* the risk-free rate, the market return and the valuation date.

From those it rebuilds holdings net of sales, cost basis at average purchase
cost, realised and unrealised P&L, money-weighted returns (XIRR), the 12-month
back-cast of the current holdings, and every risk ratio, then compares each with
the figure the workbook reports. It never reads a tick, a banner or a pass count.

Usage:

    python validate_portfolio.py                    # defaults to the workbook here
    python validate_portfolio.py path/to/Book.xlsx

Exits 0 when everything reconciles, 1 otherwise, so it can run in CI.
"""
from __future__ import annotations

import datetime as dt
import math
import statistics
import sys
import warnings
from collections import defaultdict
from pathlib import Path

warnings.filterwarnings("ignore", category=UserWarning)

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl")

DEFAULT = Path(__file__).parent / "Stock Portfolio.xlsx"
SHEETS = ("Ledger", "Dashboard", "Analytics", "Risk Analytics", "Price History", "Portfolio Series")

# Money is carried at full float precision, so the tolerance only absorbs
# floating-point noise. XIRR is iterative in Excel, so it gets its own tolerance.
ABS_TOL = 1e-6
REL_TOL = 1e-9
IRR_TOL = 1e-6
TRADING_DAYS = 252

LEDGER_FIRST_ROW = 3
DASH_FIRST, DASH_LAST = 7, 22          # Dashboard holdings rows
STOCK_FIRST, STOCK_LAST = 10, 25       # Analytics holdings rows
SECTOR_FIRST, SECTOR_LAST = 29, 38
HHI_ROW = 39
SPARK_FIRST = 3                        # Price History rows 3..18, same order as the Dashboard

# Risk Analytics, column C
RA = {"rf": "C6", "rm": "C7", "erp": "C8", "beta": "C9", "capm": "C10", "alpha": "C11",
      "treynor": "C15", "ir": "C16", "te": "C17", "sharpe": "C18", "sortino": "C19",
      "calmar": "C20", "mean": "C24", "vol": "C25", "var95": "C26", "var99": "C27",
      "cvar": "C28", "downside": "C30", "hvar": "C31", "eff_pos": "C34", "spread": "C35"}
# Portfolio Series, column B
PS = {"covered": "B10", "share": "B11", "ndays": "B12", "ret12": "B13", "vol": "B14",
      "downside": "B15", "mdd": "B16", "hvar": "B17", "spy12": "B18", "beta_spy": "B19",
      "te": "B20", "active": "B21", "ir": "B22", "asof": "B24"}


class Result:
    def __init__(self, group, label, derived, reported, tol=None, note=None):
        self.group, self.label = group, label
        self.derived, self.reported = derived, reported
        self.tol, self.note = tol, note

    @property
    def ok(self):
        d, r = self.derived, self.reported
        if not isinstance(d, (int, float)) or not isinstance(r, (int, float)):
            return False
        if self.tol is not None:
            return abs(d - r) <= self.tol
        return abs(d - r) <= max(ABS_TOL, REL_TOL * max(abs(d), abs(r)))

    @property
    def detail(self):
        d, r = self.derived, self.reported
        if not isinstance(r, (int, float)):
            return f"workbook reports {r!r}, which is not a number"
        if not isinstance(d, (int, float)):
            return f"could not derive ({d!r})"
        return f"derived {d:,.10g} vs workbook {r:,.10g}, differs by {d - r:,.3g}"


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(v))
    return None


def xirr(flows):
    """Excel's XIRR: the rate r solving sum(v / (1 + r) ^ ((d - d0) / 365)) = 0."""
    flows = sorted(flows)
    d0 = flows[0][0]
    t = [(d - d0).days / 365.0 for d, _ in flows]
    v = [x for _, x in flows]

    def f(r):
        return sum(x / (1 + r) ** ti for x, ti in zip(v, t))

    lo, hi = -0.9999, 10.0
    flo, fhi = f(lo), f(hi)
    if flo * fhi > 0:
        return None
    for _ in range(300):
        mid = (lo + hi) / 2
        fm = f(mid)
        if flo * fm <= 0:
            hi, fhi = mid, fm
        else:
            lo, flo = mid, fm
    return (lo + hi) / 2


def percentile_inc(xs, p):
    """Excel PERCENTILE.INC: linear interpolation between order statistics."""
    s = sorted(xs)
    k = p * (len(s) - 1)
    f = math.floor(k)
    c = min(f + 1, len(s) - 1)
    return s[f] + (s[c] - s[f]) * (k - f)


def load(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    missing = [s for s in SHEETS if s not in wb.sheetnames]
    if missing:
        sys.exit(f"missing sheet(s) {missing} in {path.name}; sheets are {wb.sheetnames}")
    return wb


def read_ledger(ws):
    trades = []
    for r in range(LEDGER_FIRST_ROW, ws.max_row + 1):
        kind = ws.cell(r, 4).value
        if not isinstance(kind, str) or kind.strip().lower() not in ("buy", "sell"):
            continue
        trades.append({"date": as_date(ws.cell(r, 2).value), "kind": kind.strip().lower(),
                       "price": num(ws.cell(r, 5).value), "units": num(ws.cell(r, 6).value),
                       "amount": num(ws.cell(r, 7).value), "ticker": ws.cell(r, 8).value})
    return trades


def positions(trades):
    p = defaultdict(lambda: {"bu": 0.0, "ba": 0.0, "su": 0.0, "sa": 0.0, "flows": []})
    for t in trades:
        x = p[t["ticker"]]
        if t["kind"] == "buy":
            x["bu"] += t["units"]
            x["ba"] += t["amount"]
            x["flows"].append((t["date"], -t["amount"]))
        else:
            x["su"] += t["units"]
            x["sa"] += t["amount"]
            x["flows"].append((t["date"], t["amount"]))
    for x in p.values():
        x["net"] = x["bu"] - x["su"]
        x["avg"] = x["ba"] / x["bu"] if x["bu"] else 0.0
        x["cost"] = x["net"] * x["avg"]
        x["realised"] = x["sa"] - x["su"] * x["avg"]
    return p


def build(wb):
    led, dash, an = wb["Ledger"], wb["Dashboard"], wb["Analytics"]
    ra, spark, ps = wb["Risk Analytics"], wb["Price History"], wb["Portfolio Series"]
    checks = []

    def add(group, label, derived, reported, **kw):
        checks.append(Result(group, label, derived, reported, **kw))

    def worst(pairs):
        """Largest absolute gap across (derived, reported) pairs, as a single check."""
        gaps = [abs(d - r) for d, r in pairs if num(d) is not None and num(r) is not None]
        return 0.0, (max(gaps) if len(gaps) == len(pairs) else "missing value")

    # ---------------- ledger and holdings ----------------
    trades = read_ledger(led)
    pos = positions(trades)
    asof = as_date(ps[PS["asof"]].value)
    G = "Holdings from the ledger"
    add(G, f"every trade amount = price x units ({len(trades)} trades)",
        *worst([(t["price"] * t["units"], t["amount"]) for t in trades]))

    rows = []
    for dr, ar in zip(range(DASH_FIRST, DASH_LAST + 1), range(STOCK_FIRST, STOCK_LAST + 1)):
        tk = dash.cell(dr, 13).value
        if tk not in pos:
            add(G, f"holding {tk!r} on the Dashboard appears in the ledger", None, 1.0)
            continue
        x = pos[tk]
        price = num(dash.cell(dr, 5).value)
        mv = x["net"] * price if price is not None else None
        rows.append({"tk": tk, "x": x, "price": price, "mv": mv,
                     "units": num(dash.cell(dr, 4).value), "a_cost": num(an.cell(ar, 2).value),
                     "a_mv": num(an.cell(ar, 3).value), "a_gain": num(an.cell(ar, 4).value),
                     "a_ret": num(an.cell(ar, 5).value), "a_wt": num(an.cell(ar, 6).value),
                     "a_irr": num(an.cell(ar, 8).value), "a_real": num(an.cell(ar, 9).value)})
    n = len(rows)
    add(G, f"every holding's units = bought - sold ({n} holdings)",
        *worst([(r["x"]["net"], r["units"]) for r in rows]))
    add(G, "no holding has negative units", 0.0, max(0.0, -min(r["x"]["net"] for r in rows)))
    add(G, "every holding's market value = net units x price", *worst([(r["mv"], r["a_mv"]) for r in rows]))
    add(G, "every holding's cost basis = net units x average purchase cost",
        *worst([(r["x"]["cost"], r["a_cost"]) for r in rows]))
    add(G, "every holding's unrealised P&L = market value - cost basis",
        *worst([(r["mv"] - r["x"]["cost"], r["a_gain"]) for r in rows]))
    add(G, "every holding's realised P&L = proceeds - units sold x average cost",
        *worst([(r["x"]["realised"], r["a_real"]) for r in rows]))
    add(G, "every holding's return = unrealised P&L / cost basis",
        *worst([((r["mv"] - r["x"]["cost"]) / r["x"]["cost"], r["a_ret"]) for r in rows]))
    mv_total = sum(r["mv"] for r in rows)
    add(G, "every holding's weight = market value / total",
        *worst([(r["mv"] / mv_total, r["a_wt"]) for r in rows]))
    irr_gaps = []
    for r in rows:
        derived = xirr(r["x"]["flows"] + [(asof, r["mv"])])
        irr_gaps.append((derived, r["a_irr"]))
    d, g = worst(irr_gaps)
    add(G, "every holding's annual return = XIRR of its trades plus today's value", d, g, tol=IRR_TOL)

    # ---------------- portfolio totals ----------------
    G = "Portfolio totals"
    buys = sum(t["amount"] for t in trades if t["kind"] == "buy")
    sells = sum(t["amount"] for t in trades if t["kind"] == "sell")
    cost_total = sum(r["x"]["cost"] for r in rows)
    realised_total = sum(r["x"]["realised"] for r in rows)
    add(G, "total market value = sum of holdings", mv_total, num(an["B6"].value))
    add(G, "total cost basis = sum of holdings", cost_total, num(an["A6"].value))
    add(G, "unrealised P&L = market value - cost basis", mv_total - cost_total, num(an["C6"].value))
    add(G, "realised P&L = sum of holdings' realised P&L", realised_total, num(dash["P25"].value))
    add(G, "total return = (value + sale proceeds - purchases) / purchases",
        (mv_total + sells - buys) / buys, num(an["D6"].value))
    add(G, "total gain = realised + unrealised P&L", mv_total + sells - buys,
        (num(an["C6"].value) or 0) + (num(dash["P25"].value) or 0))
    flows = [(t["date"], -t["amount"] if t["kind"] == "buy" else t["amount"]) for t in trades]
    add(G, "portfolio annual return = XIRR of every trade plus today's value",
        xirr(flows + [(asof, mv_total)]), num(an["E6"].value), tol=IRR_TOL)

    # ---------------- concentration ----------------
    G = "Concentration"
    sectors = []
    for r in range(SECTOR_FIRST, SECTOR_LAST + 1):
        name, mv = an.cell(r, 1).value, num(an.cell(r, 2).value)
        if name and mv is not None:
            sectors.append({"mv": mv, "wt": num(an.cell(r, 3).value), "hhi": num(an.cell(r, 4).value)})
    add(G, "sector market values sum to the portfolio total", sum(s["mv"] for s in sectors), mv_total)
    add(G, "sector weights sum to 100%", sum(s["wt"] for s in sectors), 1.0)
    add(G, "HHI = sum of squared sector weights", sum(s["wt"] ** 2 for s in sectors),
        num(an.cell(HHI_ROW, 4).value))
    weights = [r["mv"] / mv_total for r in rows]
    add(G, "effective positions = 1 / HHI of holding weights",
        1 / sum(w * w for w in weights), num(ra[RA["eff_pos"]].value))
    add(G, "largest / smallest position ratio", max(weights) / min(weights), num(ra[RA["spread"]].value))

    # ---------------- 12-month back-cast ----------------
    G = "12-month back-cast of current holdings"
    dates = [v for v in (spark.cell(2, c).value for c in range(2, spark.max_column + 1)) if num(v) is not None]
    ndays = len(dates)
    value = [0.0] * ndays
    covered = 0
    covered_mv = 0.0
    for i, r in enumerate(rows):
        prices = [spark.cell(SPARK_FIRST + i, c).value for c in range(2, 2 + ndays)]
        if num(prices[0]) is None:
            continue
        covered += 1
        covered_mv += r["mv"]
        for k, p in enumerate(prices):
            value[k] += r["x"]["net"] * (num(p) or 0.0)
    rets = [value[k] / value[k - 1] - 1 for k in range(1, ndays)]
    reported_value = [num(ps.cell(5, c).value) for c in range(2, 2 + ndays)]
    add(G, f"daily value = net units x close, summed ({ndays} days)",
        *worst(list(zip(value, reported_value))))
    add(G, "holdings with price history", float(covered), num(ps[PS["covered"]].value))
    add(G, "share of market value covered", covered_mv / mv_total, num(ps[PS["share"]].value))
    ret12 = value[-1] / value[0] - 1
    vol = statistics.stdev(rets) * math.sqrt(TRADING_DAYS)
    downside = math.sqrt(sum(x * x for x in rets if x < 0) / len(rets)) * math.sqrt(TRADING_DAYS)
    peak, mdd = value[0], 0.0
    for v in value:
        peak = max(peak, v)
        mdd = min(mdd, v / peak - 1)
    hvar = -percentile_inc(rets, 0.05)
    add(G, "12-month return = last value / first value - 1", ret12, num(ps[PS["ret12"]].value))
    add(G, "volatility = sample stdev of daily returns x sqrt(252)", vol, num(ps[PS["vol"]].value))
    add(G, "downside deviation = sqrt(mean of squared negative returns) x sqrt(252)",
        downside, num(ps[PS["downside"]].value))
    add(G, "max drawdown = worst fall from a running peak", mdd, num(ps[PS["mdd"]].value))
    add(G, "historical VaR (95%, 1-day) = minus the 5th percentile return", hvar, num(ps[PS["hvar"]].value))

    spy = [num(ps.cell(7, c).value) for c in range(2, 2 + ndays)]
    if all(v is not None for v in spy):
        srets = [spy[k] / spy[k - 1] - 1 for k in range(1, ndays)]
        spy12 = spy[-1] / spy[0] - 1
        mr, ms = statistics.fmean(rets), statistics.fmean(srets)
        beta_spy = (sum((a - mr) * (b - ms) for a, b in zip(rets, srets))
                    / sum((b - ms) ** 2 for b in srets))
        te = statistics.stdev([a - b for a, b in zip(rets, srets)]) * math.sqrt(TRADING_DAYS)
        add(G, "SPY 12-month return", spy12, num(ps[PS["spy12"]].value))
        add(G, "beta vs SPY = slope of daily returns", beta_spy, num(ps[PS["beta_spy"]].value))
        add(G, "tracking error = stdev of (portfolio - SPY) x sqrt(252)", te, num(ps[PS["te"]].value))
        add(G, "information ratio = active return / tracking error", (ret12 - spy12) / te,
            num(ps[PS["ir"]].value))
    else:
        add(G, "SPY closes are present for every day", None, "missing")

    # ---------------- CAPM and risk ratios ----------------
    G = "CAPM and risk ratios"
    rf, rm = num(ra[RA["rf"]].value), num(ra[RA["rm"]].value)
    betas = [num(ra.cell(r, 4).value) for r in range(39, 55)]
    beta = sum(w * b for w, b in zip(weights, betas))
    capm = rf + beta * (rm - rf)
    add(G, "equity risk premium = Rm - Rf", rm - rf, num(ra[RA["erp"]].value))
    add(G, "portfolio beta = weighted average of holding betas", beta, num(ra[RA["beta"]].value))
    add(G, "CAPM expected return = Rf + beta x (Rm - Rf)", capm, num(ra[RA["capm"]].value))
    add(G, "Jensen's alpha = 12m return - CAPM expected return", ret12 - capm, num(ra[RA["alpha"]].value))
    add(G, "Treynor = (12m return - Rf) / beta", (ret12 - rf) / beta, num(ra[RA["treynor"]].value))
    add(G, "Sharpe = (12m return - Rf) / volatility", (ret12 - rf) / vol, num(ra[RA["sharpe"]].value))
    add(G, "Sortino = (12m return - Rf) / downside deviation", (ret12 - rf) / downside,
        num(ra[RA["sortino"]].value))
    add(G, "Calmar = 12m return / |max drawdown|", ret12 / abs(mdd), num(ra[RA["calmar"]].value),
        note=("A cell in error here would still leave a Validation tab that skips "
              "errors printing ALL PASS; this workbook's tab now counts errors as failures."))
    add(G, "parametric VaR 95% = CAPM return - 1.645 x volatility", capm - 1.645 * vol,
        num(ra[RA["var95"]].value))
    add(G, "parametric VaR 99% = CAPM return - 2.326 x volatility", capm - 2.326 * vol,
        num(ra[RA["var99"]].value))
    add(G, "expected shortfall 95% = CAPM return - 2.063 x volatility", capm - 2.063 * vol,
        num(ra[RA["cvar"]].value))

    meta = {"holdings": n, "sectors": len(sectors), "mv_total": mv_total, "asof": asof,
            "covered": covered, "ndays": ndays}
    return checks, meta


def main(argv):
    path = Path(argv[1]) if len(argv) > 1 else DEFAULT
    if not path.exists():
        sys.exit(f"workbook not found: {path}")

    wb = load(path)
    checks, meta = build(wb)

    print(f"Workbook: {path.name}   Valuation date: {meta['asof']}")
    print(f"Holdings: {meta['holdings']}   Sectors: {meta['sectors']}   "
          f"Market value: {meta['mv_total']:,.2f}")
    print(f"Back-cast: {meta['covered']} of {meta['holdings']} holdings have price history, "
          f"{meta['ndays']} trading days")
    print("-" * 78)

    failed = 0
    group = None
    for c in checks:
        if c.group != group:
            group = c.group
            print(f"\n  {group}")
        if c.ok:
            print(f"    PASS  {c.label}")
        else:
            failed += 1
            print(f"    FAIL  {c.label}")
            print(f"            {c.detail}")
            if c.note:
                for line in _wrap(c.note, 64):
                    print(f"            {line}")

    print("\n" + "-" * 78)
    print(f"{len(checks) - failed} of {len(checks)} figures re-derived independently and matched.")
    if failed:
        print(f"{failed} FAILED. The workbook disagrees with its own inputs on these.")
        return 1
    print("Every figure was rebuilt from the ledger, prices and inputs, and agrees with the workbook.")
    return 0


def _wrap(text, width):
    words, line, out = text.split(), "", []
    for w in words:
        if len(line) + len(w) + 1 > width:
            out.append(line)
            line = w
        else:
            line = f"{line} {w}".strip()
    if line:
        out.append(line)
    return out


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
